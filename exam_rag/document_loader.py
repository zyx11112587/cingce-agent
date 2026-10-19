from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from .models import Chunk
from .parser import clean_text, stable_id


class DocumentDependencyError(RuntimeError):
    """Raised when an optional document reader is not installed."""


@dataclass
class PageText:
    page: int
    text: str


def _infer_doc_type(source_rel: str) -> str:
    lower = source_rel.lower()
    if "政治理论" in source_rel or "politics" in lower or "政治" in source_rel:
        return "政治理论"
    if "申论" in source_rel or "shenlun" in lower or "essay" in lower:
        return "申论"
    return "文档"


def _infer_subject(source_rel: str, doc_type: str) -> str:
    for subject in (
        "政治理论",
        "申论",
        "言语理解与表达",
        "数量关系",
        "判断推理",
        "资料分析",
        "常识判断",
    ):
        if subject in source_rel:
            return subject
    return doc_type if doc_type != "文档" else "综合"


def _infer_title(source_rel: str, text: str) -> str:
    heading = re.search(r"^#{1,6}\s+(.+)$", text, re.MULTILINE)
    if heading:
        return clean_text(heading.group(1))
    first_line = next((line.strip() for line in text.splitlines() if line.strip()), "")
    return clean_text(first_line)[:80] or Path(source_rel).stem


def _read_pdf_pages(path: Path) -> list[PageText]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise DocumentDependencyError(
            "PDF support requires pypdf. Install it with: pip install pypdf"
        ) from exc

    try:
        reader = PdfReader(str(path))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                text = _read_pdf_page_with_pdfplumber(path, page_number)
            pages.append(PageText(page=page_number, text=text))
    except DocumentDependencyError:
        raise
    except Exception as exc:
        raise ValueError(f"Unable to read PDF {path}: {exc}") from exc

    if not any(page.text.strip() for page in pages):
        raise ValueError(
            f"PDF has no extractable text: {path}. "
            "It may be a scanned PDF and needs an OCR step."
        )
    return pages


def _read_pdf_page_with_pdfplumber(path: Path, page_number: int) -> str:
    try:
        import pdfplumber
    except ImportError:
        return ""
    try:
        with pdfplumber.open(str(path)) as pdf:
            return pdf.pages[page_number - 1].extract_text() or ""
    except Exception:
        return ""


def _read_docx(path: Path) -> list[PageText]:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentDependencyError(
            "DOCX support requires python-docx. "
            "Install it with: pip install python-docx"
        ) from exc

    document = Document(str(path))
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                parts.append(" | ".join(values))
    return [PageText(page=0, text="\n\n".join(parts))]


def _read_plain_text(path: Path) -> list[PageText]:
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() in {".html", ".htm"}:
        text = html.unescape(re.sub(r"<[^>]+>", "\n", text))
    return [PageText(page=0, text=text)]


def read_document_pages(path: Path) -> list[PageText]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf_pages(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix in {".md", ".markdown", ".txt", ".html", ".htm"}:
        return _read_plain_text(path)
    raise ValueError(f"Unsupported document format: {path.suffix}")


def _paragraphs(text: str) -> list[str]:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    parts = re.split(r"\n\s*\n", text)
    return [clean_text(part) for part in parts if clean_text(part)]


def _pack_paragraphs(
    paragraphs: Iterable[str],
    max_chars: int = 900,
    overlap_chars: int = 120,
) -> list[str]:
    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            while start < len(paragraph):
                chunks.append(paragraph[start:start + max_chars].strip())
                if start + max_chars >= len(paragraph):
                    break
                start += max_chars - overlap_chars
            continue
        candidate = f"{current}\n\n{paragraph}".strip()
        if current and len(candidate) > max_chars:
            chunks.append(current.strip())
            tail = current[-overlap_chars:] if overlap_chars else ""
            current = f"{tail}\n\n{paragraph}".strip()
        else:
            current = candidate
    if current.strip():
        chunks.append(current.strip())
    return chunks


def load_document(path: Path, input_dir: Path) -> list[Chunk]:
    source_rel = str(path.relative_to(input_dir)).replace("\\", "/")
    pages = read_document_pages(path)
    full_text = "\n\n".join(page.text for page in pages if page.text.strip())
    if not full_text.strip():
        raise ValueError(f"No text extracted from document: {path}")

    doc_type = _infer_doc_type(source_rel)
    subject = _infer_subject(source_rel, doc_type)
    title = _infer_title(source_rel, full_text)
    chunks: list[Chunk] = []
    chunk_index = 0
    for page in pages:
        for piece in _pack_paragraphs(_paragraphs(page.text)):
            chunk_id = stable_id(source_rel, "document", page.page, chunk_index)
            chunks.append(
                Chunk(
                    chunk_id=chunk_id,
                    qid="",
                    record_id=stable_id(source_rel, "document", page.page, chunk_index, piece[:80]),
                    chunk_type="document",
                    text=piece,
                    title=title,
                    year="",
                    region="",
                    subject=subject,
                    question_type=doc_type,
                    source_path=str(path.resolve()),
                    source_rel=source_rel,
                    source_format=path.suffix.lower().lstrip("."),
                    doc_type=doc_type,
                    page_start=page.page,
                    page_end=page.page,
                )
            )
            chunk_index += 1
    return chunks


def iter_document_files(input_dir: Path) -> Iterable[Path]:
    suffixes = {".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm"}
    for path in sorted(input_dir.rglob("*")):
        if ".git" in path.parts or not path.is_file():
            continue
        if path.suffix.lower() in suffixes:
            yield path
